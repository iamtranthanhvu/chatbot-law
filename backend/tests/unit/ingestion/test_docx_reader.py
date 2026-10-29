from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from lawqa.ingestion.docx_reader import read_paragraphs

MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _run(text: str, tag: str = "w:t") -> etree._Element:
    run = OxmlElement("w:r")
    text_element = OxmlElement(tag)
    text_element.text = text
    run.append(text_element)
    return run


def _wrapped(wrapper_tag: str, text: str, **attrs: str) -> etree._Element:
    wrapper = OxmlElement(wrapper_tag)
    for key, value in attrs.items():
        wrapper.set(qn(key), value)
    wrapper.append(_run(text))
    return wrapper


def _build_docx(path: Path) -> Path:
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    left = table.cell(0, 0).paragraphs[0]
    left.add_run("QUỐC HỘI")
    left.add_run().add_break()
    left.add_run("-------")
    table.cell(0, 1).paragraphs[0].add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    document.add_paragraph("LUẬT")
    document.add_paragraph("Điều 1. Phạm vi điều chỉnh")
    document.add_paragraph("")
    document.add_paragraph("Luật này quy định về chế độ sở hữu đất đai.")
    document.save(str(path))
    return path


def test_reads_tables_and_paragraphs_in_document_order(tmp_path: Path) -> None:
    lines = read_paragraphs(_build_docx(tmp_path / "doc.docx"))

    assert lines == [
        "QUỐC HỘI",
        "-------",
        "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
        "LUẬT",
        "Điều 1. Phạm vi điều chỉnh",
        "",
        "Luật này quy định về chế độ sở hữu đất đai.",
    ]


def test_reads_every_cell_of_a_large_table(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=6, cols=4)
    expected = []
    for r in range(6):
        for c in range(4):
            text = f"ô {r}-{c}"
            table.cell(r, c).paragraphs[0].add_run(text)
            expected.append(text)
    path = tmp_path / "large.docx"
    document.save(str(path))

    assert read_paragraphs(path) == expected


def test_merged_cells_are_read_once(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.paragraphs[0].add_run("CHỦ TỊCH QUỐC HỘI")
    path = tmp_path / "merged.docx"
    document.save(str(path))

    assert read_paragraphs(path) == ["CHỦ TỊCH QUỐC HỘI"]


def test_reads_text_inside_smart_tags_and_other_run_wrappers(tmp_path: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Cộng hòa xã hội chủ nghĩa Việt ")
    paragraph._p.append(
        _wrapped(
            "w:smartTag",
            "Nam",
            **{
                "w:uri": "urn:schemas-microsoft-com:office:smarttags",
                "w:element": "country-region",
            },
        )
    )
    paragraph._p.append(_run("; người "))
    paragraph._p.append(_wrapped("w:customXml", "Việt Nam", **{"w:element": "place"}))
    paragraph._p.append(_wrapped("w:ins", " định cư", **{"w:id": "1", "w:author": "a"}))
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "2")
    deleted.set(qn("w:author"), "a")
    deleted.append(_run(" CHỮ ĐÃ XÓA", tag="w:delText"))
    paragraph._p.append(deleted)
    path = tmp_path / "wrapped.docx"
    document.save(str(path))

    assert read_paragraphs(path) == ["Cộng hòa xã hội chủ nghĩa Việt Nam; người Việt Nam định cư"]


def test_alternate_content_is_read_once(tmp_path: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Điều 1. ")
    alternate = etree.SubElement(paragraph._p, f"{{{MC_NS}}}AlternateContent")
    choice = etree.SubElement(alternate, f"{{{MC_NS}}}Choice", Requires="wps")
    choice.append(_run("Phạm vi"))
    fallback = etree.SubElement(alternate, f"{{{MC_NS}}}Fallback")
    fallback.append(_run("Phạm vi"))
    path = tmp_path / "alternate.docx"
    document.save(str(path))

    assert read_paragraphs(path) == ["Điều 1. Phạm vi"]
